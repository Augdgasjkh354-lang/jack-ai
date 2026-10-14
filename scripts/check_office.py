"""用 python-docx / python-pptx / openpyxl 打开单元测试生成的样例，确认文件格式没问题。"""
import sys
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

d = sys.argv[1]
doc = Document(f"{d}/t.docx")
print("docx paragraphs:", [p.text for p in doc.paragraphs][:8], "tables:", len(doc.tables))
assert doc.paragraphs[0].text.startswith("2026") and len(doc.tables) == 1

prs = Presentation(f"{d}/t.pptx")
for i, s in enumerate(prs.slides):
    print("slide", i + 1, [sh.text_frame.text for sh in s.shapes if sh.has_text_frame])
assert len(prs.slides) == 4
assert any(sh.shape_type == 13 for sh in prs.slides[3].shapes), "slide 4 should have a picture"

wb = load_workbook(f"{d}/t.xlsx")
for ws in wb.worksheets:
    print("sheet", ws.title, [[c.value for c in r] for r in ws.iter_rows()])
assert wb.worksheets[0]["D2"].value == "=B2*C2" and wb.worksheets[0]["B2"].value == 3.5
assert wb.worksheets[0]["B3"].value == "007"
print("OK")
